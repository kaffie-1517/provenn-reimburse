"""Excel export of approved reimbursements, for import into an ERP."""

import io
from collections.abc import Iterable
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import Verification
from app.refcode import display

COLUMNS = [
    ("Submitted", 18),
    ("Employee", 22),
    ("Email", 28),
    ("Reference", 14),
    ("Vendor", 28),
    ("Invoice date", 13),
    ("Amount", 14),
    ("Currency", 9),
    ("Check", 10),
    ("Approved", 18),
    ("Approved by", 28),
]


def _naive(dt: datetime | None) -> datetime | None:
    # Excel has no time zones; write UTC wall-clock time.
    return dt.replace(tzinfo=None) if dt else None


def approved_workbook(rows: Iterable[Verification], approver_names: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Approved"
    ws.append([name for name, _ in COLUMNS])

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4A4237")
    for i, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = head, fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    for v in rows:
        inv = v.invoice
        ws.append(
            [
                _naive(v.submitted_at),
                v.submitter.name,
                v.submitter.email,
                display(inv.reference_code) if inv else None,
                inv.vendor_name if inv else None,
                inv.invoice_date if inv else None,
                inv.amount_cents / 100 if inv else None,
                inv.currency if inv else None,
                v.result,
                _naive(v.approved_at),
                approver_names.get(v.approved_by),
            ]
        )
        r = ws.max_row
        ws.cell(r, 1).number_format = ws.cell(r, 10).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(r, 6).number_format = "yyyy-mm-dd"
        ws.cell(r, 7).number_format = "#,##0.00"

    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
