import io

from openpyxl import load_workbook

from app import worker
from tests.pdfs import invoice_pdf, tamper
from tests.test_invoices import issue


async def issued_pdf(client, auth, storage) -> tuple[str, bytes]:
    """Issues an invoice as a provider and returns (code, the downloaded PDF)."""
    code = (await issue(client, await auth.provider()))["reference_code"]
    await worker.drain(storage)
    r = await client.get(f"/api/v1/invoices/{code}/download")
    return code, r.content


async def submit(client, headers, data: bytes, name: str = "claim.pdf"):
    r = await client.post(
        "/api/v1/verifications", files={"pdf": (name, data, "application/pdf")}, headers=headers
    )
    assert r.status_code == 201, r.text
    return r.json()


async def test_match_mismatch_not_found(client, auth, storage):
    code, original = await issued_pdf(client, auth, storage)
    _, emp = await auth.company("acme")

    ok = await submit(client, emp, original)
    assert ok["result"] == "match"
    assert ok["approval_status"] == "pending"
    assert ok["invoice"]["reference_code"] == code
    assert ok["invoice"]["amount_cents"] == 150000
    assert ok["submitter_email"] == "emp@acme.example.com"

    forged = await submit(client, emp, tamper(original))
    assert forged["result"] == "mismatch"
    assert forged["extracted_code"] == code

    unknown = await submit(client, emp, invoice_pdf(vendor="Never issued"))
    assert unknown["result"] == "not_found"
    assert unknown["invoice"] is None


async def test_listing_is_scoped(client, auth, storage):
    _, original = await issued_pdf(client, auth, storage)
    acme_admin, acme_emp = await auth.company("acme")
    beta_admin, beta_emp = await auth.company("beta")
    await submit(client, acme_emp, original)
    await submit(client, acme_emp, tamper(original))
    await submit(client, beta_emp, original)

    async def listing(h, **params):
        r = await client.get("/api/v1/verifications", headers=h, params=params)
        assert r.status_code == 200, r.text
        return r.json()

    assert (await listing(acme_emp))["total"] == 2
    assert (await listing(acme_admin))["total"] == 2
    assert (await listing(beta_admin))["total"] == 1
    mism = await listing(acme_admin, result="mismatch")
    assert [v["result"] for v in mism["items"]] == ["mismatch"]
    page = await listing(acme_admin, limit=1)
    assert page["total"] == 2 and len(page["items"]) == 1


async def test_approval_rules(client, auth, storage):
    _, original = await issued_pdf(client, auth, storage)
    acme_admin, acme_emp = await auth.company("acme")
    beta_admin, _ = await auth.company("beta")
    v = await submit(client, acme_emp, original)
    url = f"/api/v1/verifications/{v['id']}"

    # Other company's admin can't see it at all.
    r = await client.patch(url, json={"decision": "approved"}, headers=beta_admin)
    assert r.status_code == 404
    # Employees can't decide.
    r = await client.patch(url, json={"decision": "approved"}, headers=acme_emp)
    assert r.status_code == 403
    r = await client.patch(url, json={"decision": "maybe"}, headers=acme_admin)
    assert r.status_code == 422

    r = await client.patch(url, json={"decision": "approved"}, headers=acme_admin)
    assert r.status_code == 200
    assert r.json()["approval_status"] == "approved"
    assert r.json()["result"] == "match"  # untouched by the decision

    r = await client.patch(url, json={"decision": "rejected"}, headers=acme_admin)
    assert r.status_code == 409


async def test_export_contains_only_approved_rows(client, auth, storage):
    code, original = await issued_pdf(client, auth, storage)
    acme_admin, acme_emp = await auth.company("acme")
    beta_admin, beta_emp = await auth.company("beta")

    keep = await submit(client, acme_emp, original)
    reject = await submit(client, acme_emp, tamper(original))
    await submit(client, acme_emp, original)  # stays pending
    other = await submit(client, beta_emp, original)

    for vid, decision, h in [
        (keep["id"], "approved", acme_admin),
        (reject["id"], "rejected", acme_admin),
        (other["id"], "approved", beta_admin),
    ]:
        r = await client.patch(
            f"/api/v1/verifications/{vid}", json={"decision": decision}, headers=h
        )
        assert r.status_code == 200

    r = await client.get("/api/v1/verifications/export", headers=acme_admin)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][1] == "emp@acme.example.com"
    assert rows[0][2] == f"PNN-{code}"
    assert rows[0][5] == 1500.0
    assert rows[0][9] == "admin@acme.example.com"

    assert (await client.get("/api/v1/verifications/export", headers=acme_emp)).status_code == 403


async def test_only_employees_submit(client, auth, storage):
    _, original = await issued_pdf(client, auth, storage)
    admin, _ = await auth.company("acme")
    provider = await auth.provider("p2@example.com")
    for h in (admin, provider):
        r = await client.post(
            "/api/v1/verifications",
            files={"pdf": ("c.pdf", original, "application/pdf")},
            headers=h,
        )
        assert r.status_code == 403
